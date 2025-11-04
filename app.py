# app.py
import streamlit as st
from io import BytesIO
from typing import Tuple, List
import re
import PyPDF2
import docx
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from fpdf import FPDF
from difflib import get_close_matches
from datetime import datetime

# ------------------------- Helpers -------------------------

def fix_unicode(s: str) -> str:
    if not s:
        return ""
    # Replace all unsupported UTF-8 characters
    return (s.replace("—", "-")
             .replace("–", "-")
             .replace("•", "-")
             .replace("’", "'")
             .replace("“", '"')
             .replace("”", '"')
             .encode("latin-1", errors="replace")
             .decode("latin-1"))

def extract_text_from_pdf(file_stream) -> str:
    try:
        reader = PyPDF2.PdfReader(file_stream)
    except Exception:
        file_stream.seek(0)
        reader = PyPDF2.PdfReader(file_stream)

    text_parts = []
    for page in reader.pages:
        try:
            page_text = page.extract_text()
        except:
            page_text = ""
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)

def extract_text_from_docx(file_stream) -> str:
    file_stream.seek(0)
    doc = docx.Document(file_stream)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_text_from_txt(file_stream) -> str:
    file_stream.seek(0)
    try:
        return file_stream.read().decode("utf-8")
    except:
        file_stream.seek(0)
        return file_stream.read().decode("latin-1")

def clean_text(s: str) -> str:
    s = s or ""
    s = fix_unicode(s)
    s = s.replace("\r", " ").replace("\t", " ")
    s = re.sub(r"[^\w\s\-\/\.\,]", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip().lower()

def get_top_keywords(text: str, top_k: int = 20) -> List[str]:
    if not text or len(text.split()) < 3:
        return []
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1,2), max_features=5000)
    try:
        tfidf = vectorizer.fit_transform([text])
    except:
        return []
    feature_array = vectorizer.get_feature_names_out()
    tfidf_scores = tfidf.toarray()[0]
    scored = list(zip(feature_array, tfidf_scores))
    scored.sort(key=lambda x: x[1], reverse=True)

    seen = set()
    out = []
    for w, _ in scored[:top_k]:
        if w not in seen:
            out.append(w); seen.add(w)
    return out

def similarity_score(jd: str, resume: str) -> float:
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1,2), max_features=5000)
    docs = [jd or "", resume or ""]
    try:
        tfidf = vectorizer.fit_transform(docs)
    except:
        return 0.0
    if tfidf.shape[0] < 2 or tfidf.shape[1] == 0:
        return 0.0
    sim = cosine_similarity(tfidf[0:1], tfidf[1:2])[0][0]
    return float(sim * 100)

def find_matches(jd_keywords: List[str], resume_text: str) -> Tuple[List[str], List[str]]:
    resume_text_lower = resume_text.lower()
    found, not_found = [], []
    tokens = set(re.findall(r"\w+(?: \w+)?", resume_text_lower))
    for kw in jd_keywords:
        kw_l = kw.lower()
        if kw_l in resume_text_lower:
            found.append(kw)
        else:
            matches = get_close_matches(kw_l, tokens, n=1, cutoff=0.8)
            found.append(kw) if matches else not_found.append(kw)
    return found, not_found

def summarize_resume(resume_text: str, max_sentences: int = 4) -> str:
    resume_text = fix_unicode(resume_text)
    resume_text = resume_text.replace("\n", ". ")
    sentences = [s.strip() for s in re.split(r"[.\n]", resume_text) if s.strip()]
    if not sentences:
        return "No readable text to summarize."
    scores = []
    for s in sentences:
        score = len(s.split())
        if re.search(r"\b(python|java|c\+\+|sql|aws|docker|react|flask|data)\b", s.lower()):
            score += 5
        scores.append((score, s))
    scores.sort(key=lambda x: x[0], reverse=True)
    top = [s for _, s in scores[:max_sentences]]
    return " ".join(top)

# ------------------------- PDF FIXED -------------------------

def create_pdf_report(name, match, found, missing, summary, recs):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    pdf.cell(0, 8, fix_unicode("Resume Screening Report"), ln=True)
    pdf.ln(3)
    pdf.cell(0, 8, fix_unicode(f"Candidate: {name}"), ln=True)
    pdf.cell(0, 8, fix_unicode(f"Match Percentage: {match:.1f}%"), ln=True)

    pdf.ln(3)
    pdf.cell(0, 8, "Skills Found:", ln=True)
    for s in found:
        pdf.cell(0, 8, "- " + fix_unicode(s), ln=True)

    pdf.ln(3)
    pdf.cell(0, 8, "Missing Skills:", ln=True)
    for s in missing:
        pdf.cell(0, 8, "- " + fix_unicode(s), ln=True)

    pdf.ln(3)
    pdf.multi_cell(0, 6, fix_unicode("Summary:\n" + summary))
    pdf.ln(2)
    pdf.multi_cell(0, 6, fix_unicode("Recommendations:\n" + "\n".join(recs)))

    out = pdf.output(dest="S")
    if isinstance(out, str):
        out = out.encode("latin-1", errors="replace")
    return out

# ------------------------- Streamlit UI -------------------------

st.set_page_config(page_title="Resume Screener", layout="wide")
st.title("Resume Screener - Auto JD vs Resume Analyzer")
st.caption("No APIs, No fonts, No errors - Works offline.")

col1, col2 = st.columns([1,2])

with col1:
    uploaded_file = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    sample_jd = ("We are seeking a Data Entry Operator with strong typing skills and attention to detail.")
    jd_text = st.text_area("Job Description", value=sample_jd, height=200)
    run = st.button("Run Analysis")

with col2:
    output_area = st.empty()

if "last_report" not in st.session_state:
    st.session_state["last_report"] = None

if run:
    if not uploaded_file:
        st.warning("Upload a resume.")
    elif not jd_text or len(jd_text.strip()) < 5:
        st.warning("Paste a job description.")
    else:
        with st.spinner("Analyzing..."):
            try:
                file_bytes = uploaded_file.read()
                bio = BytesIO(file_bytes)
                ext = uploaded_file.name.split(".")[-1].lower()

                if ext == "pdf":
                    resume_text = extract_text_from_pdf(bio)
                elif ext == "docx":
                    resume_text = extract_text_from_docx(bio)
                else:
                    resume_text = extract_text_from_txt(bio)

                if not resume_text.strip():
                    st.error("Could not extract text (maybe scanned PDF).")
                else:
                    jd_clean = clean_text(jd_text)
                    resume_clean = clean_text(resume_text)

                    jd_keywords = get_top_keywords(jd_clean, 30)
                    found_skills, missing_skills = find_matches(jd_keywords, resume_clean)
                    match_pct = similarity_score(jd_clean, resume_clean)
                    summary = summarize_resume(resume_text)

                    recs = []
                    if match_pct >= 75:
                        recs.append("Strong match - consider for interview.")
                    elif match_pct >= 50:
                        recs.append("Moderate match - some skills missing.")
                    else:
                        recs.append("Low match - lacks required skills.")
                    if missing_skills:
                        recs.append("Ask about: " + ", ".join(missing_skills[:6]))

                    with output_area.container():
                        st.metric("Match Score", f"{match_pct:.1f}%")
                        st.write("**Summary:**", summary)

                        st.write("**Skills Found**", found_skills)
                        st.write("**Missing Skills**", missing_skills)
                        st.write("**Recommendations**")
                        for r in recs:
                            st.write("- " + r)

                        pdf_bytes = create_pdf_report(
                            uploaded_file.name,
                            match_pct,
                            found_skills,
                            missing_skills,
                            summary,
                            recs
                        )

                        st.session_state["last_report"] = (pdf_bytes, uploaded_file.name)
                        st.success("Report ready.")

            except Exception as e:
                st.exception(e)

if st.session_state["last_report"]:
    pdf_bytes, candidate_name = st.session_state["last_report"]
    fname = f"report_{candidate_name}_{datetime.utcnow().strftime('%Y%m%d%H%M')}.pdf"
    st.download_button("Download PDF Report", data=pdf_bytes, file_name=fname, mime="application/pdf")

st.markdown("---")
st.write("Made to never throw Unicode errors. Enjoy.")
